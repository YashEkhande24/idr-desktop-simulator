import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation('SIH2026-IDEA-Presentation-Format.pptx')
print(f"Slide width: {prs.slide_width.inches} in, height: {prs.slide_height.inches} in")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for s in slide.shapes:
        info = f"ID: {s.shape_id}, Name: '{s.name}', Type: {s.shape_type}"
        if s.has_text_frame:
            text = s.text_frame.text.replace('\n', ' // ')
            info += f" | Text: {text[:100]}"
        print(info)
