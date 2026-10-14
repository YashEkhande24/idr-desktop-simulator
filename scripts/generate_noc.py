import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), b_val.get('val', 'single'))
            node.set(qn('w:sz'), str(b_val.get('sz', 4)))
            node.set(qn('w:space'), '0')
            node.set(qn('w:color'), b_val.get('color', 'auto'))
            tcBorders.append(node)
        else:
            node = OxmlElement(f'w:{b_name}')
            node.set(qn('w:val'), 'none')
            tcBorders.append(node)
    tcPr.append(tcBorders)

def create_noc_document(output_docx):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
        
    # Institutional Header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(1)
    r0 = p_inst.add_run("Amrutvahini Sheti & Shikshan Vikas Sanstha's\n")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(9.5)
    
    r1 = p_inst.add_run("AMRUTVAHINI COLLEGE OF ENGINEERING, SANGAMNER\n")
    r1.font.name = "Times New Roman"
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = RGBColor(16, 44, 87)
    
    r2 = p_inst.add_run("(An Autonomous Institute Affiliated to Savitribai Phule Pune University, Pune)\nApproved by AICTE, New Delhi | Accredited 'A+' Grade by NAAC | NBA Accredited Programs\nPost: Sangamner S.K. (422608), Tal: Sangamner, Dist: Ahmednagar, Maharashtra, India\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(70, 70, 70)
    
    r3 = p_inst.add_run("DEPARTMENT OF ELECTRONICS & TELECOMMUNICATION ENGINEERING")
    r3.font.name = "Times New Roman"
    r3.font.bold = True
    r3.font.size = Pt(10.5)
    r3.font.color.rgb = RGBColor(16, 44, 87)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(4)
    r_div = p_div.add_run("―" * 58)
    r_div.font.size = Pt(9)
    r_div.font.color.rgb = RGBColor(150, 150, 150)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ref & Date
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_after = Pt(6)
    r_ref = p_ref.add_run("Ref. No.: AVCOE/ETC/2026-27/CR-NOC/01")
    r_ref.font.name = "Times New Roman"
    r_ref.font.bold = True
    r_ref.font.size = Pt(10)
    
    r_tab = p_ref.add_run("\t\t\t\t\t\t")
    r_date = p_ref.add_run("Date: 09/10/2026")
    r_date.font.name = "Times New Roman"
    r_date.font.bold = True
    r_date.font.size = Pt(10)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("NO OBJECTION CERTIFICATE (NOC)")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.underline = True
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("(Issued for Copyright Registration under Section 2(o) of The Copyright Act, 1957 & Rule 70 of Copyright Rules, 2013)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(8.5)
    r_sub.font.italic = True

    # Cert text
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_after = Pt(4)
    r_c = p_cert.add_run("This is to certify that the following students are bonafide undergraduate scholars of Amrutvahini College of Engineering, Sangamner, affiliated to Savitribai Phule Pune University, Pune:")
    r_c.font.name = "Times New Roman"
    r_c.font.size = Pt(9.5)

    # Student table
    students = [
        ("1", "Mr. Yash Kailas Ekhande", "Electronics & Telecommunication Engg.", "Undergraduate Scholar (Lead)"),
        ("2", "Mr. Abhijeet Dilip Gholap", "Electronics & Telecommunication Engg.", "Undergraduate Scholar"),
        ("3", "Ms. Ankita Prabhakar Kadu", "Electronics & Telecommunication Engg.", "Undergraduate Scholar"),
        ("4", "Mr. Krushna Sanjay Shinde", "Automation & Robotics Engineering", "Undergraduate Scholar"),
        ("5", "Ms. Shreya Gopal Dahale", "Computer Engineering", "Undergraduate Scholar"),
        ("6", "Ms. Kanchan Vijay Jadhav", "Automation & Robotics Engineering", "Undergraduate Scholar")
    ]
    
    tbl_students = doc.add_table(rows=len(students)+1, cols=4)
    tbl_students.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_students.autofit = False
    
    headers = ["Sr. No.", "Name of Author / Applicant", "Department / Discipline", "Academic Status"]
    widths = [Inches(0.6), Inches(2.3), Inches(2.3), Inches(1.8)]
    
    for c_idx, h_text in enumerate(headers):
        cell = tbl_students.cell(0, c_idx)
        cell.width = widths[c_idx]
        set_cell_background(cell, "102C57")
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.font.name = "Times New Roman"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, (s_no, s_name, s_dept, s_status) in enumerate(students):
        row_cells = tbl_students.rows[r_idx+1].cells
        bg = "F4F7FB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text_val in enumerate([s_no, s_name, s_dept, s_status]):
            c = row_cells[c_idx]
            c.width = widths[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 40, 40, 80, 80)
            set_cell_borders(c, 
                             top={'val': 'single', 'sz': 4, 'color': 'D0D7DE'},
                             bottom={'val': 'single', 'sz': 4, 'color': 'D0D7DE'},
                             left={'val': 'none'}, right={'val': 'none'})
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text_val)
            r.font.name = "Times New Roman"
            if c_idx == 1:
                r.font.bold = True
            r.font.size = Pt(8.5)

    # Work description
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(6)
    p_desc.paragraph_format.space_after = Pt(3)
    p_desc.paragraph_format.line_spacing = 1.15
    r_d1 = p_desc.add_run("The above-named individuals have jointly conceptualized, developed, and authored the original software codebase, algorithmic architecture, and technical documentation titled:")
    r_d1.font.name = "Times New Roman"
    r_d1.font.size = Pt(9.5)

    # Title box
    tbl_tbox = doc.add_table(rows=1, cols=1)
    tbl_tbox.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_t = tbl_tbox.cell(0, 0)
    c_t.width = Inches(7.0)
    set_cell_background(c_t, "F4F6F9")
    set_cell_margins(c_t, 60, 60, 100, 100)
    set_cell_borders(c_t,
                     top={'val': 'single', 'sz': 6, 'color': 'B0B8C4'},
                     bottom={'val': 'single', 'sz': 6, 'color': 'B0B8C4'},
                     left={'val': 'single', 'sz': 6, 'color': 'B0B8C4'},
                     right={'val': 'single', 'sz': 6, 'color': 'B0B8C4'})
    p_tb = c_t.paragraphs[0]
    p_tb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tt = p_tb.add_run("“AI-ML Based Intelligent Dead Reckoning (IDR) System for Seamless Vehicle Navigation in GNSS-Denied Environments”\n")
    r_tt.font.name = "Times New Roman"
    r_tt.font.bold = True
    r_tt.font.size = Pt(10)
    r_tt.font.color.rgb = RGBColor(16, 44, 87)
    r_subt = p_tb.add_run("(Software Artifact: IDR-Navigator | Team Inventor, SIH Problem ID: SIH26168)")
    r_subt.font.name = "Times New Roman"
    r_subt.font.size = Pt(8.5)
    r_subt.font.italic = True

    # Guidance and Clearance
    p_clear = doc.add_paragraph()
    p_clear.paragraph_format.space_before = Pt(5)
    p_clear.paragraph_format.space_after = Pt(4)
    p_clear.paragraph_format.line_spacing = 1.15
    r_c1 = p_clear.add_run("under the academic guidance and research supervision of ")
    r_c1.font.name = "Times New Roman"
    r_c1.font.size = Pt(9.5)
    r_c2 = p_clear.add_run("Prof. Dr. R. P. Labade")
    r_c2.font.name = "Times New Roman"
    r_c2.font.bold = True
    r_c2.font.size = Pt(9.5)
    r_c3 = p_clear.add_run(", Head of the Department, Department of Electronics & Telecommunication Engineering, Amrutvahini College of Engineering, Sangamner.\n\n")
    r_c3.font.name = "Times New Roman"
    r_c3.font.size = Pt(9.5)
    
    r_c4 = p_clear.add_run("Amrutvahini College of Engineering, Sangamner and the Department of Electronics & Telecommunication Engineering have ")
    r_c4.font.name = "Times New Roman"
    r_c4.font.size = Pt(9.5)
    r_c5 = p_clear.add_run("NO OBJECTION")
    r_c5.font.name = "Times New Roman"
    r_c5.font.bold = True
    r_c5.font.size = Pt(9.5)
    r_c6 = p_clear.add_run(" whatsoever to the filing, processing, and registration of the copyright for the said computer software, algorithmic framework, and literary documentation in the joint names of the declared authors and supervisor with the ")
    r_c6.font.name = "Times New Roman"
    r_c6.font.size = Pt(9.5)
    r_c7 = p_clear.add_run("Copyright Office, Government of India (ROC, New Delhi)")
    r_c7.font.name = "Times New Roman"
    r_c7.font.bold = True
    r_c7.font.size = Pt(9.5)
    r_c8 = p_clear.add_run(". The Institution affirms that this is an original academic research output and waives any claims or objections against the copyright registration.")
    r_c8.font.name = "Times New Roman"
    r_c8.font.size = Pt(9.5)

    # Signature block
    tbl_sig = doc.add_table(rows=2, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.autofit = False
    
    sig_widths = [Inches(2.3), Inches(2.3), Inches(2.3)]
    for row in tbl_sig.rows:
        for idx, c in enumerate(row.cells):
            c.width = sig_widths[idx]
            set_cell_borders(c)
            
    # Row 0: Signature spacing
    for c in tbl_sig.rows[0].cells:
        p = c.paragraphs[0]
        p.add_run("\n\n")
        
    # Row 1: Names and Designations
    sig_data = [
        ("Mr. V. K. Abhang", "SPOC, Smart India Hackathon\nAssistant Professor, Dept. of E&TC\nAmrutvahini College of Engg."),
        ("Prof. Dr. R. P. Labade", "Head of the Department\nDept. of E&TC Engineering\nAmrutvahini College of Engg."),
        ("Dr. P. R. Baviskar", "Principal\nAmrutvahini College of Engineering\nSangamner, SPPU Pune")
    ]
    
    for c_idx, (name, desig) in enumerate(sig_data):
        cell = tbl_sig.cell(1, c_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rn = p.add_run(name + "\n")
        rn.font.name = "Times New Roman"
        rn.font.bold = True
        rn.font.size = Pt(9.5)
        rd = p.add_run(desig)
        rd.font.name = "Times New Roman"
        rd.font.size = Pt(8)
        rd.font.color.rgb = RGBColor(60, 60, 60)

    # Official stamp placeholder
    p_stamp = doc.add_paragraph()
    p_stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stamp.paragraph_format.space_before = Pt(6)
    r_st = p_stamp.add_run("[ Official Institutional Seal / Departmental Stamp ]")
    r_st.font.name = "Times New Roman"
    r_st.font.size = Pt(7.5)
    r_st.font.italic = True
    r_st.font.color.rgb = RGBColor(120, 120, 120)

    doc.save(output_docx)
    print(f"Created clean 1-page NOC docx: {output_docx}")

if __name__ == '__main__':
    create_noc_document(r'e:\inventor\NOC Copyright.docx')
