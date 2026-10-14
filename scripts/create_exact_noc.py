import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_no_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def create_noc_files():
    doc = docx.Document()
    
    # 1 inch margins all around
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    
    # Top Ref and Date line matching exact sample
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_after = Pt(28)
    
    r_ref = p_ref.add_run("Ref. No. AVCOE/2024-25/ \t\t\t\t\t\tDate: 30/09/2024")
    r_ref.font.name = "Times New Roman"
    r_ref.font.size = Pt(12)
    
    # Title matching exact sample
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(32)
    
    r_title = p_title.add_run("No Objection Certificate")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    
    # Body paragraph with exact template sentence structure
    p_body = doc.add_paragraph()
    p_body.paragraph_format.line_spacing = 1.5
    p_body.paragraph_format.space_after = Pt(72)
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r1 = p_body.add_run("This is to certify that ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    
    # Student names (bolded or plain - let's make the filled names clear)
    students_text = "Mr. Yash Kailas Ekhande, Mr. Abhijeet Dilip Gholap, Ms. Ankita Prabhakar Kadu, Mr. Krushna Sanjay Shinde, Ms. Shreya Gopal Dahale, Ms. Kanchan Vijay Jadhav"
    r_names = p_body.add_run(students_text)
    r_names.font.name = "Times New Roman"
    r_names.font.size = Pt(12)
    r_names.font.bold = True
    
    r2 = p_body.add_run(" is students of Department of ")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    
    r_dept1 = p_body.add_run("Electronics & Telecommunication")
    r_dept1.font.name = "Times New Roman"
    r_dept1.font.size = Pt(12)
    r_dept1.font.bold = True
    
    r3 = p_body.add_run(" engineering at Amrutvahini college of Engineering, Sangamner , Savitribai Phule Pune University, Pune  and  topic is “")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    
    r_topic = p_body.add_run("AI-ML Based Intelligent Dead Reckoning (IDR) System for Seamless Vehicle Navigation in GNSS-Denied Environments")
    r_topic.font.name = "Times New Roman"
    r_topic.font.size = Pt(12)
    r_topic.font.bold = True
    
    r4 = p_body.add_run("” under the supervision of ")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(12)
    
    r_guide = p_body.add_run("Dr. R. P. Labade")
    r_guide.font.name = "Times New Roman"
    r_guide.font.size = Pt(12)
    r_guide.font.bold = True
    
    r5 = p_body.add_run(", Department of ")
    r5.font.name = "Times New Roman"
    r5.font.size = Pt(12)
    
    r_dept2 = p_body.add_run("Electronics & Telecommunication")
    r_dept2.font.name = "Times New Roman"
    r_dept2.font.size = Pt(12)
    r_dept2.font.bold = True
    
    r6 = p_body.add_run(" engineering at Amrutvahini college of Engineering, Sangamner , Savitribai Phule Pune University, Pune . This is issued for the purpose of copyright of research topic.")
    r6.font.name = "Times New Roman"
    r6.font.size = Pt(12)
    
    # Signatures: 3 columns matching original sample layout exactly
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(2.2), Inches(2.2), Inches(2.2)]
    
    for row in table.rows:
        for i, w in enumerate(col_widths):
            cell = row.cells[i]
            cell.width = w
            set_no_borders(cell)
            set_cell_margins(cell, 0, 0, 0, 0)
            
    # Row 0: Signature vertical space
    for c in table.rows[0].cells:
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.add_run("\n\n\n")
        
    # Row 1: Signatories
    sig_data = [
        ("Mr. V.K. Abhang", "SPOC SIH"),
        ("Dr. R.P. Labade", "Head of the Dept"),
        ("Dr. P.R.Baviskar", "Principal")
    ]
    
    for i, (name, desig) in enumerate(sig_data):
        p = table.rows[1].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        rn = p.add_run(name + "\n")
        rn.font.name = "Times New Roman"
        rn.font.size = Pt(12)
        rn.font.bold = True
        
        rd = p.add_run(desig)
        rd.font.name = "Times New Roman"
        rd.font.size = Pt(11)

    docx_path = r"e:\inventor\NOC Copyright.docx"
    doc.save(docx_path)
    print(f"Saved: {docx_path}")

if __name__ == '__main__':
    create_noc_files()
